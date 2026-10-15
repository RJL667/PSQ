#!/usr/bin/env python3
"""gen_questionnaire.py — build questionnaire_model.json from the rating workbook + proposal form.

The workbook's "Cybersecurity Questionaire" sheet is the scoring source of truth: its H column
scores each answer (weights in C/D/E/F), H282 is the 0..1 cyber posture that the price sheet bands
into the maturity multiplier (C11 -> C12), and its I column carries the underwriting flags.

This script copies those cells VERBATIM (formulas as text, constants as values) plus every
external lookup cell they reference, so questionnaire.py can re-evaluate the sheet for any set of
answers without Excel. It also maps the fillable proposal-form PDF's fields onto questionnaire rows.

Usage:  python3 tools/gen_questionnaire.py [workbook.xlsx] [proposal_form.pdf]
Defaults look in ../_source/Corporate Rating Engine/ (gitignored client working files).
"""
import difflib
import glob
import json
import os
import re
import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
SVC = HERE.parent
SRC = SVC / "_source" / "Corporate Rating Engine"
OUT = SVC / "questionnaire_model.json"


def _find(pattern):
    hits = [p for p in glob.glob(str(SRC / pattern)) if not os.path.basename(p).startswith("~$")]
    return hits[0] if hits else None


XLSX = sys.argv[1] if len(sys.argv) > 1 else _find("*Sec Questions*.xlsx")
PDF = sys.argv[2] if len(sys.argv) > 2 else _find("*V2 2026*.pdf")
if not XLSX:
    sys.exit("workbook not found (pass its path)")

REF_RE = re.compile(r"(?:'([^']+)'|([A-Za-z_][A-Za-z0-9_.]*))!\$?([A-Z]{1,3})\$?(\d+)(?::\$?([A-Z]{1,3})\$?(\d+))?")
BULLETS = "\u00b7\u2022\u25e6-\u00a0 \t"


def col_idx(letters):
    n = 0
    for ch in letters:
        n = n * 26 + (ord(ch) - 64)
    return n


def col_letters(n):
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def plain(v):
    if v is None:
        return None
    if hasattr(v, "text") and not isinstance(v, str):  # ArrayFormula
        v = v.text
    if isinstance(v, (int, float, bool, str)):
        return v
    return str(v)


def norm(s):
    s = (s or "").lower().replace("\u2019", "'").replace("\u00a0", " ")
    s = s.lstrip(BULLETS)
    s = re.sub(r"[^a-z0-9 ]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def key(s):
    return re.sub(r"[^a-z0-9]", "", (s or "").lower())


wbf = openpyxl.load_workbook(XLSX, data_only=False)
wbv = openpyxl.load_workbook(XLSX, data_only=True)
QS = next(s for s in wbf.sheetnames if "question" in s.lower())
PS = next(s for s in wbf.sheetnames
          if isinstance(wbf[s]["C11"].value, str) and "H282" in wbf[s]["C11"].value)
qf, qv = wbf[QS], wbv[QS]

# ---- questionnaire cells (A..R), verbatim ---------------------------------
sheets = {QS: {}, PS: {}}
for row in qf.iter_rows(min_row=1, max_row=qf.max_row, max_col=18):
    for c in row:
        v = plain(c.value)
        if v is None or v == "":
            continue
        sheets[QS][c.coordinate] = v
for coord in ("C11", "C12", "D12"):
    sheets[PS][coord] = plain(wbf[PS][coord].value)

# ---- external cells referenced by those formulas (lookup scales, maturity bands) ----
for sheet, cells in list(sheets.items()):
    for v in list(cells.values()):
        if not (isinstance(v, str) and v.startswith("=")):
            continue
        for m in REF_RE.finditer(v):
            ref_sheet = m.group(1) or m.group(2)
            if ref_sheet == sheet:
                continue
            c1, r1 = col_idx(m.group(3)), int(m.group(4))
            c2, r2 = (col_idx(m.group(5)), int(m.group(6))) if m.group(5) else (c1, r1)
            dst = sheets.setdefault(ref_sheet, {})
            for r in range(r1, r2 + 1):
                for c in range(c1, c2 + 1):
                    coord = "%s%d" % (col_letters(c), r)
                    if coord in dst:
                        # Already copied verbatim (e.g. the price sheet's C11 -> questionnaire H282).
                        # Never overwrite a formula with its cached value, or the posture freezes.
                        continue
                    val = plain(wbv[ref_sheet][coord].value)
                    if val is not None:
                        dst[coord] = val

# ---- cached results (for validating the evaluator against Excel) ----------
cached = {QS: {}, PS: {}}
for r in range(30, 284):
    for L in ("H", "I"):
        v = qv["%s%d" % (L, r)].value
        if isinstance(v, (int, float, str)) and not isinstance(v, bool) and (sheets[QS].get("%s%d" % (L, r)) or "").__class__ is str:
            if str(sheets[QS].get("%s%d" % (L, r), "")).startswith("="):
                cached[QS]["%s%d" % (L, r)] = v
for coord in ("C11", "C12", "D12"):
    cached[PS][coord] = plain(wbv[PS][coord].value)

# ---- question rows ----------------------------------------------------------
formula_text = " ".join(v for v in sheets[QS].values() if isinstance(v, str) and v.startswith("="))
scored_rows = sorted({int(n) for n in re.findall(r"(?<![A-Z$])\$?G\$?(\d+)", formula_text)})
flag_rows = sorted(int(c[1:]) for c, v in sheets[QS].items()
                   if c.startswith("I") and c[1:].isdigit() and isinstance(v, str)
                   and re.search(r"G\d+\s*=", v))

section = subsection = ""
questions = []
prev_plain = None   # last non-bullet question row
prev_root = None    # last non-bullet, non-"If ..." question row
for r in range(4, 282):
    text = sheets[QS].get("B%d" % r)
    if not isinstance(text, str) or not text.strip():
        continue
    t = text.strip()
    if re.match(r"^Section\s+\d+", t):
        section, subsection = t, ""
        continue
    if re.match(r"^[A-Z]\.\s", t) and "?" not in t:
        subsection = t
        continue
    opts = [plain(qf.cell(r, c).value) for c in range(10, 19)]
    opts = [str(o).strip() for o in opts if o not in (None, "") and str(o).strip()]
    stripped = t.lstrip(BULLETS)
    low = stripped.lower()
    is_bullet = t[:1] in "\u00b7\u2022\u25e6"
    is_if = low.startswith(("if yes", "if no", "if parent"))
    yesno = [o.lower() for o in opts[:2]] == ["yes", "no"]
    parent = gate = None
    if is_bullet and prev_plain:
        parent = prev_plain
    elif is_if and prev_root:
        parent = prev_root
    if parent:
        pq = next((q for q in questions if q["row"] == parent), None)
        if pq and pq["yesno"]:
            gate = "No" if low.startswith("if no") else "Yes"
        else:
            parent = None
    q = {"row": r, "text": stripped, "section": section, "subsection": subsection,
         "options": opts, "yesno": yesno, "scored": r in scored_rows,
         "flag": r in flag_rows or sheets[QS].get("A%d" % r) == "*",
         "parent": parent, "gate": gate, "default": plain(qv["G%d" % r].value)}
    questions.append(q)
    if not is_bullet:
        prev_plain = r
        if not is_if:
            prev_root = r

QBY = {q["row"]: q for q in questions}

# ---- proposal form: map fillable fields onto questionnaire rows ------------
form = {"pdf": None, "fields": {}, "details": {}, "unmapped": [], "name_conflicts": []}
if PDF:
    import fitz  # PyMuPDF

    form["pdf"] = os.path.basename(PDF)
    doc = fitz.open(PDF)
    cand = [(q["row"], norm(q["text"])) for q in questions]

    def match(label):
        n = norm(label)
        if len(n) < 6:
            return [], 0.0
        scored = [(difflib.SequenceMatcher(None, n, c).ratio(), r) for r, c in cand]
        best = max(s for s, _ in scored)
        return sorted(r for s, r in scored if s >= best - 0.005), best

    instances = []  # one per question line on the form
    for page in doc:
        lines = []
        for b in page.get_text("dict")["blocks"]:
            for ln in b.get("lines", []):
                txt = "".join(sp["text"] for sp in ln["spans"]).strip()
                if txt:
                    lines.append((*ln["bbox"], txt))
        for wd in (page.widgets() or []):
            if wd.field_type_string == "Signature":
                continue
            rc = wd.rect
            cy = (rc.y0 + rc.y1) / 2
            same = [l for l in lines if l[1] - 4 <= cy <= l[3] + 4]
            qs_ = [l for l in same if len(l[4]) > 12 and l[0] < rc.x0]
            label = min(qs_, key=lambda l: l[0])[4] if qs_ else ""
            if not label:
                above = [l for l in lines if l[3] <= rc.y0 + 2 and rc.y0 - l[3] < 18
                         and len(l[4]) > 12 and l[0] < rc.x0]
                if above:
                    label = max(above, key=lambda l: l[3])[4]
            right = sorted([l for l in same if l[0] >= rc.x1 - 3], key=lambda l: l[0] - rc.x1)
            w = {"name": wd.field_name, "type": wd.field_type_string, "x": round(rc.x0, 1),
                 "y": round(rc.y0, 1), "page": page.number + 1, "label": label,
                 "near": right[0][4][:30] if right else ""}
            if wd.field_type_string == "CheckBox":
                try:
                    w["on"] = str(wd.on_state())
                except Exception:
                    w["on"] = "Yes"
            if wd.field_type_string == "ComboBox":
                w["choices"] = [str(c[0] if isinstance(c, (list, tuple)) else c) for c in (wd.choice_values or [])]
            inst = next((i for i in instances if i["page"] == w["page"] and i["label"] == label
                         and label and abs(i["y"] - w["y"]) < 6), None)
            if inst is None:
                inst = {"page": w["page"], "y": w["y"], "label": label, "widgets": []}
                instances.append(inst)
            inst["widgets"].append(w)

    # assign instances to rows; identical-text rows are filled in form order
    groups = {}
    for inst in instances:
        rows, score = match(inst["label"])
        inst["score"] = round(score, 3)
        if score >= 0.75 and rows:
            groups.setdefault(tuple(rows), []).append(inst)
        else:
            inst["row"] = None
    for rows, insts in groups.items():
        insts.sort(key=lambda i: (i["page"], i["y"]))
        for k, inst in enumerate(insts):
            inst["row"] = rows[k] if k < len(rows) else None

    # option-name fallback: unmatched checkboxes named after an option (frameworks, SPF/DKIM/DMARC)
    optkeys = {}
    for q in questions:
        if not q["yesno"]:
            for o in q["options"]:
                optkeys.setdefault(key(o), q["row"])
    loose = []
    for inst in instances:
        for w in inst["widgets"]:
            if w["type"] == "CheckBox" and (inst["row"] is None or len(inst["widgets"]) > 2) and key(w["name"]) in optkeys:
                loose.append((optkeys[key(w["name"])], w))
            elif inst["row"] is None:
                form["unmapped"].append({k: w[k] for k in ("page", "type", "name", "label")})
    for r, w in loose:
        f = form["fields"].setdefault(str(r), {"kind": "multi", "fields": {}})
        f.setdefault("fields", {})[w["name"]] = w["name"]

    def yn_suffix(name):
        m = re.search(r"(?:^|[\s_]|(?<=[a-z]))([YyNn])$", name.strip())
        return None if not m else ("Yes" if m.group(1).lower() == "y" else "No")

    for inst in instances:
        r = inst["row"]
        if r is None:
            continue
        ws_ = inst["widgets"]
        boxes = [w for w in ws_ if w["type"] == "CheckBox"]
        combos = [w for w in ws_ if w["type"] == "ComboBox"]
        texts = [w for w in ws_ if w["type"] == "Text"]
        q = QBY[r]
        if combos:
            form["fields"][str(r)] = {"kind": "choice", "field": combos[0]["name"],
                                      "choices": combos[0].get("choices", [])}
        elif len(boxes) == 2 and q["yesno"]:
            yes, no = sorted(boxes, key=lambda w: w["x"])
            form["fields"][str(r)] = {"kind": "yn", "yes": yes["name"], "no": no["name"],
                                      "on": [yes.get("on", "Yes"), no.get("on", "Yes")]}
            sy, sn = yn_suffix(yes["name"]), yn_suffix(no["name"])
            if sy == "No" or sn == "Yes":
                form["name_conflicts"].append({"row": r, "yes": yes["name"], "no": no["name"]})
        elif boxes and str(r) not in form["fields"]:
            form["fields"][str(r)] = {"kind": "multi",
                                      "fields": {w["name"]: (w["near"] if re.match(r"(?i)check ?box", w["name"]) else w["name"])
                                                 for w in boxes},
                                      "on": {w["name"]: w.get("on", "Yes") for w in boxes}}
        for w in texts:
            form["details"].setdefault(str(r), []).append(w["name"])

    form["instances"] = [{"page": i["page"], "y": i["y"], "row": i["row"], "score": i["score"],
                          "label": i["label"][:80], "widgets": [w["name"] for w in i["widgets"]]}
                         for i in instances]

model = {
    "generated_from": {"xlsx": os.path.basename(XLSX), "pdf": form["pdf"]},
    "questionnaire_sheet": QS,
    "price_sheet": PS,
    "posture_cell": "H282",
    "flag_total_cell": "I282",
    "flag_expected_cell": "J282",
    "sheets": sheets,
    "cached": cached,
    "questions": questions,
    "form": form,
}
OUT.write_text(json.dumps(model, ensure_ascii=False, indent=1, default=str), encoding="utf-8")

mapped = {int(k) for k in form["fields"]}
report = {
    "out": str(OUT), "questions": len(questions), "scored_rows": len(scored_rows),
    "flag_rows": flag_rows, "external_sheets": {k: len(v) for k, v in sheets.items()},
    "cached_cells": len(cached[QS]),
    "form_fields_mapped": len(mapped),
    "scored_rows_without_form_field": [(r, QBY[r]["text"][:60]) for r in scored_rows
                                       if r in QBY and r not in mapped],
    "name_conflicts": form["name_conflicts"],
    "unmapped_widgets": len(form["unmapped"]),
}
print(json.dumps(report, ensure_ascii=False, indent=1, default=str))
